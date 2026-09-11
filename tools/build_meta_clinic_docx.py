from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "commercial"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "1F4E79"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Arial", size=10.5, bold=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, space_before, space_after in (
        ("Title", 24, 0, 10),
        ("Heading 1", 16, 15, 6),
        ("Heading 2", 12.5, 11, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Meta Clinic")
    set_run_font(footer_run, size=8.5, color=RGBColor(90, 90, 90))


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    set_run_font(run, size=24, bold=True)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    run = sub.add_run(subtitle)
    set_run_font(run, size=11, color=RGBColor(80, 80, 80))


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_properties.append(header_marker)
    header_cells = table.rows[0].cells
    for index, value in enumerate(headers):
        cell = header_cells[index]
        if widths:
            cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_run_font(run, size=9, bold=True, color=RGBColor(255, 255, 255))

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            if widths:
                cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_run_font(run, size=8.6)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    label_run = p.add_run(label + ": ")
    set_run_font(label_run, bold=True, size=9)
    url_run = p.add_run(url)
    set_run_font(url_run, size=9, color=RGBColor(31, 78, 121))


def pricing_strategy():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Meta Clinic Pricing and Packaging Strategy"
    doc.core_properties.author = "Meta Clinic"
    add_title(doc, "Meta Clinic Pricing and Packaging Strategy", "Future-state mature platform pricing hypothesis")

    add_body(doc, "This pricing view is for the fully completed product, not the current prototype or an immediate customer quote. The mature product assumed here is a complete ambulatory operating platform for NP- and PA-led cash-pay clinics, including clinical charting, scheduling, patient access, billing and checkout, memberships/packages, retail-pharmacy e-prescribing, supplier purchasing, inventory control and supported third-party integrations. Examples such as routing prescriptions to Walgreens and connecting to suppliers are future-state assumptions, not current production-readiness claims.")
    add_body(doc, "The commercial model is software plus FDE implementation: a recurring per-clinic subscription and a separate fixed-scope fee for onboarding, approved integrations and migration. A general consultancy or done-for-you clinic-launch service is not part of this offer.")

    doc.add_heading("Product position", level=1)
    add_body(doc, "Tebra is a direct comparator. Its public offering combines configurable practice management, EHR, billing and patient-engagement capabilities. Customizability is table stakes, not the main differentiator. Meta Clinic should earn a premium by combining a focused cash-pay ambulatory workflow with FDE-led migration, retail-pharmacy prescribing, supplier purchasing and inventory operations in one coherent workspace.")
    add_body(doc, "The mature value story is time-to-operational-use and reduced system sprawl. The current build's capabilities and local verification do not establish that the future-state pharmacy, supplier, clinical, security, professional or release gates are complete.")

    doc.add_heading("Recommended model", level=1)
    add_numbered(doc, [
        "Recurring software subscription priced per clinic, with provider and site bands and ordinary administrative users included.",
        "Separate fixed-scope onboarding and FDE integration or migration fee.",
        "Optional managed configuration/support after handoff, with written hours and response targets.",
        "Vendor pass-through or separately stated charges for external services not controlled by Meta Clinic.",
    ])

    doc.add_heading("Package ladder", level=1)
    add_table(doc, ["Package", "What the buyer gets", "Price hypothesis"], [
        ("Practice Core", "Scheduling, patient access, encounters, configurable charting, supported retail-pharmacy e-prescribing, workspaces, permissions, cash-pay checkout, invoices, receipts and basic reporting.", "$799-$1,299 per clinic/month plus $5,000-$10,000 FDE onboarding"),
        ("Practice Growth", "Core plus memberships/packages, patient portal, reminders, refills/recalls, lab coordination, advanced billing, supplier purchasing, inventory controls, reorder workflows and analytics.", "$1,499-$2,499 per clinic/month plus $10,000-$20,000 FDE onboarding"),
        ("Practice Pro / Multi-site", "Growth plus multi-site controls, advanced purchasing/inventory, supplier integrations, cross-site reporting, priority support and implementation governance.", "$2,500-$4,500 per clinic/month or site-group/month plus $20,000-$40,000 FDE onboarding"),
        ("Enterprise / Network", "Contracted network scale, enterprise reporting, rollout governance, dedicated FDE capacity and negotiated service levels.", "$5,000+/month; implementation generally $40,000+"),
        ("FDE Integration and Migration", "Source discovery, mapping, approved pharmacy/supplier/data integration work, dry runs, reconciliation, validation, cutover and handoff.", "$5,000-$10,000 simple; $10,000-$20,000 standard; $20,000-$40,000+ complex"),
    ], widths=[1.45, 3.75, 1.7])

    doc.add_heading("Pricing alternatives", level=1)
    add_table(doc, ["Strategy", "Strength", "Risk and recommendation"], [
        ("Flat per clinic", "Simple and aligned with the operating unit.", "Recommended starting point; add provider/site bands later."),
        ("Per provider", "Familiar to EHR buyers.", "Penalizes growth and ignores operational staff; use only if provider support drives most cost."),
        ("Modular add-ons", "Lets a clinic expand gradually.", "Keep to two meaningful tiers plus FDE rather than feature-by-feature billing."),
        ("Usage-based", "May align price to activity.", "Defer until usage, reconciliation and billing evidence exists."),
        ("FDE retainer", "Monetizes configuration and integration expertise.", "Use written hours, response limits and change orders to prevent unbounded labor."),
    ], widths=[1.45, 2.2, 3.25])

    doc.add_heading("Commercial guardrails", level=1)
    add_bullets(doc, [
        "Keep software, onboarding and FDE deliverables distinct in the agreement and invoice.",
        "Treat real patient or clinical data migration as a gated implementation activity, not a routine import promise.",
        "Keep custom integrations, complex migration, professional services and vendor costs outside the base subscription unless explicitly contracted.",
        "Do not use unverified compliance, clinical-approval, prescribing-readiness or vendor-availability claims in pricing or sales material.",
        "Define provider, site, integration, inventory-location, support-hour, response-time, export and cancellation limits.",
        "Offer annual prepay discounts only after support cost, churn risk and cash needs are known.",
    ])

    doc.add_heading("Validation plan", level=1)
    add_table(doc, ["Test", "Price shown", "Evidence to capture"], [
        ("Mature platform test", "$799-$2,499 per clinic/month", "Buyer type, replacement value, objection, paid commitment, preferred tier and expected value."),
        ("FDE test", "$5,000-$10,000 simple and $10,000-$20,000 standard", "Source systems, data volume, mapping effort, validation hours, rework and change requests."),
        ("Feature readiness test", "Synthetic prescribing and purchasing workflows", "Reliability, vendor costs, support load and evidence required before publishing claims."),
    ], widths=[1.45, 2.0, 3.45])
    add_body(doc, "Interest is not a sale. A deposit or signed scoped engagement is stronger evidence than a positive conversation.")

    doc.add_heading("What not to promise yet", level=1)
    add_bullets(doc, [
        "A fully accepted production EHR before the relevant clinical, integration, security, professional and release gates are met.",
        "Universal clinical workflow coverage, live Walgreens routing or supplier connectivity without contract and acceptance evidence.",
        "Automatic migration of every source record without source review and reconciliation.",
        "Compliance certification, clinical approval, professional coverage or unrestricted real-data access.",
    ])

    doc.add_heading("Market references", level=1)
    add_body(doc, "These public prices are directional anchors, not an apples-to-apples valuation. Jane lists U.S. plans at $54, $79 and $99 per month. CharmHealth lists a free plan, an encounter plan and a $200 per provider/month plan. Tebra publishes a broader $49-$799 per provider/month range depending on provider type and configuration.")
    add_source(doc, "Jane pricing", "https://jane.app/pricing?locale=en-US")
    add_source(doc, "CharmHealth pricing", "https://www.charmhealth.com/ehr/ehr-pricing-us.html")
    add_source(doc, "Tebra pricing", "https://www.tebra.com/pricing/overview")

    path = OUT / "Meta_Clinic_Pricing_and_Packaging_Strategy.docx"
    doc.save(path)
    return path


def fde_offer():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Meta Clinic FDE Integration Offer"
    doc.core_properties.author = "Meta Clinic"
    add_title(doc, "Meta Clinic FDE Integration Offer", "Data migration and implementation scope")

    add_body(doc, "The FDE offer helps a clinic move from its current systems into the fully completed Meta Clinic platform through a controlled, documented process. The FDE configures the clinic, maps approved data, coordinates pharmacy, supplier or other integration work, validates the result with clinic staff and hands the configured workspace to clinic ownership. The FDE is an implementation role, not a general consultancy, clinician or clinic administrator.")

    doc.add_heading("FDE outcome", level=1)
    add_body(doc, "At the end of an accepted engagement, the clinic should have a configured Meta Clinic workspace, a documented source-to-target mapping, a reconciled migration result, an identified exception list and a staff-approved handoff. The outcome is a validated implementation within an agreed scope; it is not a guarantee of clinical, legal, regulatory or professional approval.")

    doc.add_heading("Engagement phases", level=1)
    add_table(doc, ["Phase", "FDE responsibilities", "Exit evidence"], [
        ("Discover", "Inventory current systems, exports, integrations, data domains, owners and constraints.", "Signed scope, source inventory, access plan and migration assumptions."),
        ("Map", "Define field mapping, transformations, exclusions, duplicate handling and target ownership.", "Versioned mapping and exception rules approved by the clinic."),
        ("Build", "Configure clinic settings and create the approved import or integration path.", "Configured workspace, integration notes and repeatable import procedure."),
        ("Validate", "Run a dry run, reconcile counts and samples, surface defects and guide staff review.", "Reconciliation results, exception list and staff validation record."),
        ("Handoff", "Complete agreed cutover steps, document support boundaries and train clinic operators.", "Acceptance signoff, runbook, open-items list and support transition."),
    ], widths=[1.0, 3.8, 2.1])

    doc.add_heading("Integration tiers", level=1)
    add_table(doc, ["Tier", "Typical scope", "Price hypothesis"], [
        ("Simple", "One source, clean export, standard mapping and one validated import.", "$5,000-$10,000"),
        ("Standard", "Multiple sources or custom mapping, reconciliation, staff validation and guided cutover.", "$10,000-$20,000"),
        ("Complex", "Custom pharmacy/supplier APIs, multiple sites, incomplete or high-volume records, or exceptional rollback needs.", "$20,000-$40,000+"),
    ], widths=[1.0, 4.25, 1.65])
    add_body(doc, "The fee is based on source count, data volume, mapping complexity, integration method, validation effort, cutover risk and expected FDE hours. It is not based on patient-care revenue.")

    doc.add_heading("Included and excluded work", level=1)
    add_table(doc, ["Included when scoped", "Not included by default"], [
        ("Source-system discovery and approved export review", "Legal, clinical, accounting or professional advice"),
        ("Field mapping and transformation rules", "Credential approval, clinical decisions or prescribing authority"),
        ("Approved pharmacy, supplier and inventory integration coordination when contracted", "Live pharmacy/supplier access without approved contracts, credentials and acceptance"),
        ("Configuration of supported Meta Clinic settings", "Unlimited custom development or unsupported integrations"),
        ("Dry runs, reconciliation and staff validation", "Correction of source records outside the agreed process"),
        ("Cutover runbook, handoff and defined post-cutover support", "Vendor fees, source export charges, travel or extra support unless contracted"),
    ], widths=[3.35, 3.55])

    doc.add_heading("Data and access boundary", level=1)
    add_body(doc, "The current implementation-specialist role is intentionally limited to facilities configuration. It does not grant patient, encounter, note, medication, billing, inventory or role-management authority. If a live migration requires additional access, use a separately approved migration identity or controlled process with least privilege, time bounds, auditability and revocation. Do not expand the ordinary FDE role by assumption.")
    add_body(doc, "Real patient or clinical data may be accessed or moved only after documented customer authorization, applicable agreements, source-system permissions, data-flow and security review, approved credentials, mapping, reconciliation, rollback planning and release approval. Synthetic development data is not a substitute for customer migration evidence.")

    doc.add_heading("Migration acceptance", level=1)
    add_numbered(doc, [
        "The customer confirms the source systems, authorized data domains and records to be migrated.",
        "The FDE and customer approve the versioned mapping, transformations, exclusions and exception rules.",
        "A dry run produces counts, reconciliation results and a list of unresolved exceptions.",
        "Authorized clinic staff review representative results and confirm that the target workflow is usable.",
        "The cutover uses a documented rollback or stop path and records the final outcome.",
        "The customer accepts the configured workspace, runbook, open items and support transition.",
    ])

    doc.add_heading("Commercial terms", level=1)
    add_bullets(doc, [
        "Bill the FDE engagement separately from the recurring software subscription.",
        "Use a fixed scope with a defined source count, data domains, environments, dry runs, validation sessions and handoff deliverables.",
        "Use change orders for new sources, new data domains, custom transformations, repeated failed exports, custom APIs or additional cutover attempts.",
        "Define payment milestones, customer prerequisites, cancellation, rollback, support duration and data export/exit handling.",
    ])

    doc.add_heading("Customer prerequisites", level=1)
    add_bullets(doc, [
        "A named customer owner who can approve scope, mapping and acceptance.",
        "Authorized access to source exports or APIs and confirmation that the clinic may use them for migration.",
        "A source-system inventory with vendor contacts, export limits and known data-quality issues.",
        "Clinic staff available for validation and workflow acceptance.",
        "A confirmed target environment and support contact before cutover.",
    ])

    doc.add_heading("Positioning", level=1)
    add_body(doc, "Tebra and other broad platforms already offer configurability. The FDE offer should therefore be positioned around reducing migration friction and time to a validated operating workspace for small cash-pay ambulatory clinics. The proof point is repeatable implementation evidence: shorter configuration time, fewer unresolved exceptions, clear ownership and a clean handoff.")

    path = OUT / "Meta_Clinic_FDE_Integration_Offer.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for output in (pricing_strategy(), fde_offer()):
        print(output)
